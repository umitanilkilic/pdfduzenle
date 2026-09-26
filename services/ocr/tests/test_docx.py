import io

from docx import Document

from app.docx import markdown_to_docx

MD = """## Yıllık Rapor

Şirketimiz **2026** yılında büyüdü.
Aynı paragraf.

- Birinci madde
- İkinci madde

| Ay | Gelir |
|---|---|
| Ocak | 10 |
"""


def test_converts_headings_paragraphs_lists_and_tables():
    doc = Document(io.BytesIO(markdown_to_docx(MD)))
    paras = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert paras[0] == ("Heading 2", "Yıllık Rapor")
    assert paras[1] == ("Normal", "Şirketimiz 2026 yılında büyüdü. Aynı paragraf.")
    assert [t for s, t in paras if s == "List Bullet"] == ["Birinci madde", "İkinci madde"]
    bold = [r.text for r in doc.paragraphs[1].runs if r.bold]
    assert bold == ["2026"]
    [table] = doc.tables
    assert [[c.text for c in row.cells] for row in table.rows] == [["Ay", "Gelir"], ["Ocak", "10"]]


def test_plain_text_becomes_paragraphs():
    doc = Document(io.BytesIO(markdown_to_docx("satır bir\n\nsatır iki\n")))
    assert [p.text for p in doc.paragraphs] == ["satır bir", "satır iki"]
