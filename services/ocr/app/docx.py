"""Minimal Markdown → DOCX conversion for OCR output (headings, paragraphs, lists, tables)."""

from __future__ import annotations

import io
import re

from docx import Document

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}")
_EMPHASIS = re.compile(r"\*\*(.+?)\*\*|__(.+?)__")


def markdown_to_docx(markdown: str) -> bytes:
    doc = Document()
    lines = markdown.splitlines()
    i = 0
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            _add_runs(doc.add_paragraph(), " ".join(paragraph))
            paragraph.clear()

    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            flush()
        elif line.lstrip().startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                if not _TABLE_SEPARATOR.match(lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            _add_table(doc, rows)
            continue
        elif m := _HEADING.match(line):
            flush()
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif m := _BULLET.match(line):
            flush()
            _add_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := _NUMBERED.match(line):
            flush()
            _add_runs(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            paragraph.append(line.strip())
        i += 1
    flush()

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _add_runs(p, text: str) -> None:
    """Adds text to a paragraph, turning **bold** spans into bold runs."""
    pos = 0
    for m in _EMPHASIS.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos : m.start()])
        p.add_run(m.group(1) or m.group(2)).bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])


def _add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
